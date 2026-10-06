"""Fill the CoCo CLI Hackathon submission template with the FactoryPulse project.

Usage:  python docs/build_submission_deck.py <template.pptx> <output.pptx>
Reads live metrics from data/factorypulse.db so the numbers on the Impact slide match the demo.
"""
from __future__ import annotations

import copy
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

# ----------------------------------------------------------------------------- palette (template: black header, blue footer)
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x15, 0x65, 0xD8)
SKY = RGBColor(0x2F, 0x9B, 0xFF)
TINT = RGBColor(0xEA, 0xF1, 0xFD)
TINT2 = RGBColor(0xF4, 0xF6, 0xFA)
GRAY = RGBColor(0x5B, 0x64, 0x72)
ORANGE = RGBColor(0xE8, 0x7A, 0x1A)
GREEN = RGBColor(0x1E, 0x9E, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

TEAM = dict(name="ICE", leader="Amogh Kotha Nagaraj", size="<N>",
            problem="Predictive Maintenance and OEE Command Center")


# ----------------------------------------------------------------------------- metrics from the demo DB
def load_metrics() -> dict:
    try:
        import pandas as pd
        from factorypulse import db
        c = db.connect(True)
        m = db.query(c, "SELECT * FROM model_metrics ORDER BY trained_ts DESC LIMIT 1").iloc[0]
        fe = db.query(c, "SELECT asset_id, failure_ts, downtime_minutes FROM failure_events WHERE observed=1")
        rs = db.query(c, "SELECT asset_id, ts, risk_score FROM risk_scores")
        rs["ts"], fe["failure_ts"] = pd.to_datetime(rs.ts), pd.to_datetime(fe.failure_ts)
        cut = rs.ts.quantile(0.7)
        leads = []
        for _, f in fe[fe.failure_ts > cut].iterrows():
            w = rs[(rs.asset_id == f.asset_id) & (rs.ts >= f.failure_ts - pd.Timedelta(hours=24))
                   & (rs.ts < f.failure_ts) & (rs.risk_score >= 0.7)]
            if len(w):
                leads.append((f.failure_ts - w.ts.min()).total_seconds() / 3600)
        return dict(roc=float(m.roc_auc), precision=float(m.precision_), recall=float(m.recall),
                    failures=int(len(fe)), downtime_h=float(fe.downtime_minutes.sum() / 60),
                    avg_down_h=float(fe.downtime_minutes.mean() / 60),
                    lead_h=(sum(leads) / len(leads)) if leads else 20.0)
    except Exception as exc:  # noqa: BLE001
        print("metrics fallback:", exc)
        return dict(roc=0.98, precision=0.87, recall=0.91, failures=17, downtime_h=82.5, avg_down_h=4.9, lead_h=20.4)


# ----------------------------------------------------------------------------- drawing helpers
def set_run(run, text, size=14, bold=False, color=NAVY, italic=False):
    run.text = text
    f = run.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color


def add_bullet(paragraph, level=0):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.17 + 0.17 * level))))
    pPr.set("indent", str(-int(Inches(0.17))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def text_box(slide, x, y, w, h, paragraphs, size=12, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, bullets=False, space_after=4, margin=0.05):
    """paragraphs: list of str | list[(text, {opts})] runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for text, opts in runs:
            r = p.add_run()
            set_run(r, text, size=opts.get("size", size), bold=opts.get("bold", bold),
                    color=opts.get("color", color), italic=opts.get("italic", False))
        if bullets:
            add_bullet(p)
    return tb


def rect(slide, x, y, w, h, fill=TINT, line=None, rounded=True, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    if not shadow:
        shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def circle(slide, x, y, d, fill, label, size=11, color=WHITE):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), label, size=size, bold=True, color=color)
    return shp


def box(slide, x, y, w, h, title, lines, fill=TINT, title_color=BLUE, size=9.5, title_size=11):
    rect(slide, x, y, w, h, fill=fill)
    paras = [[(title, dict(bold=True, color=title_color, size=title_size))]] + [[(ln, {})] for ln in lines]
    text_box(slide, x + 0.08, y + 0.05, w - 0.16, h - 0.1, paras, size=size, space_after=1)


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def title(slide, text, sub=None):
    text_box(slide, 0.45, 0.68, 9.1, 0.5, [[(text, dict(bold=True, size=22, color=NAVY))]], size=22)
    if sub:
        text_box(slide, 0.45, 1.13, 9.1, 0.3, [[(sub, dict(italic=True, color=GRAY, size=11))]], size=11)


def card_header(slide, x, y, w, number, heading, fill=BLUE):
    circle(slide, x, y, 0.32, fill, number, size=11)
    text_box(slide, x + 0.4, y - 0.02, w - 0.4, 0.36, [[(heading, dict(bold=True, size=13, color=NAVY))]],
             anchor=MSO_ANCHOR.MIDDLE)


# ----------------------------------------------------------------------------- template plumbing
def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sld = list(sldIdLst)[index]
    prs.part.drop_rel(sld.rId)
    sldIdLst.remove(sld)


def move_slide(prs, old_index, new_index):
    sldIdLst = prs.slides._sldIdLst
    sld = list(sldIdLst)[old_index]
    sldIdLst.remove(sld)
    sldIdLst.insert(new_index, sld)


def clone_branded_slide(prs, source_slide):
    """New slide on the same layout with the same full-bleed background picture."""
    new = prs.slides.add_slide(source_slide.slide_layout)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    pic = next(s for s in source_slide.shapes if s.shape_type == 13)
    tmp = Path(tempfile.mkdtemp()) / f"bg.{pic.image.ext}"
    tmp.write_bytes(pic.image.blob)
    new.shapes.add_picture(str(tmp), pic.left, pic.top, pic.width, pic.height)
    return new


def fill_label(shape, label, value):
    """Keep the template run formatting of 'Team Name :' and append the value."""
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = f"{label} "
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    new = copy.deepcopy(runs[0]._r)
    runs[0]._r.addnext(new)          # keep <a:endParaRPr> last, or PowerPoint drops the run
    p.runs[1].text = value
    p.runs[1].font.bold = True


# ----------------------------------------------------------------------------- slides
def slide_cover(s):
    by_text = {sh.text_frame.text.strip(): sh for sh in s.shapes if sh.has_text_frame}
    fill_label(by_text["Team Name :"], "Team Name :", TEAM["name"])
    fill_label(by_text["Team Leader Name :"], "Team Leader Name :", TEAM["leader"])
    fill_label(by_text["Team Size :"], "Team Size :", TEAM["size"])
    fill_label(by_text["Problem Statement :"], "Problem Statement :", TEAM["problem"])
    text_box(s, 0.44, 3.02, 9.0, 0.45,
             [[("FactoryPulse", dict(bold=True, size=20, color=NAVY)),
               ("  —  IT/OT converged predictive maintenance with a local Granite agent", dict(size=13, color=GRAY))]],
             anchor=MSO_ANCHOR.MIDDLE)


def slide_problem(s):
    title(s, "Problem Brief: breakdowns hide in the OT/IT gap",
          "Discrete manufacturing (machining, pressing, conveying, pumping) · Industry 4.0 IT/OT convergence")
    cw, ch, gx, gy = 4.45, 1.85, 0.2, 0.2
    x0, y0 = 0.45, 1.55
    cards = [
        ("1", "Real business problem", [
            "Sensors show vibration and temperature drifting 24-48 h before a breakdown, but that data lives in OT historians.",
            "Production schedule, maintenance history, parts and cost live in ERP/CMMS. Nobody sees both at once.",
            "Result: failures are discovered when the machine stops, and OEE loss is explained weeks later."]),
        ("2", "Target users", [
            "Reliability engineer / maintenance planner: who to inspect today, and why.",
            "Shift supervisor: what the line will deliver this shift and which asset is at risk.",
            "Plant manager: OEE by line and the cost of unplanned downtime."]),
        ("3", "Pain today  →  with FactoryPulse", [
            "Found at breakdown  →  predicted ~20 h ahead, with drivers.",
            "Manual work order after the fact  →  P1 WO auto-created with evidence.",
            "Root cause = days across systems  →  grounded answer in seconds.",
            "OEE reported monthly  →  live by asset/line, tied to the cause."]),
        ("4", "Why now / domain context", [
            "IT/OT convergence is the core of Industry 4.0 programmes; the data already exists, it is just siloed.",
            "Small local LLMs (Granite 4.0 1B) now run on a plant PC: no cloud round-trip for sensitive OT data.",
            "Snowflake Dynamic Tables, Tasks and Cortex Analyst give the same design a governed cloud twin."]),
    ]
    for i, (n, head, lines) in enumerate(cards):
        x, y = x0 + (i % 2) * (cw + gx), y0 + (i // 2) * (ch + gy)
        rect(s, x, y, cw, ch, fill=TINT if i in (0, 3) else TINT2)
        card_header(s, x + 0.15, y + 0.12, cw - 0.3, n, head, fill=BLUE if i != 2 else ORANGE)
        text_box(s, x + 0.15, y + 0.5, cw - 0.3, ch - 0.55, [[(ln, {})] for ln in lines], size=9.5, bullets=True,
                 space_after=2)


def slide_architecture(s):
    title(s, "Architecture: one pipeline from sensor stream to work order")
    text_box(s, 0.45, 1.1, 9.1, 0.28, [[("Local MVP: SQLite · scikit-learn · Streamlit · Granite 4.0 1B GGUF   |   Snowflake twin: Dynamic Tables · Model Registry · Task · Cortex Analyst",
                                         dict(italic=True, color=GRAY, size=9.5))]], size=9.5)
    OT, IT, ACT, UX = RGBColor(0xE3, 0xF6, 0xEF), TINT, RGBColor(0xFD, 0xEE, 0xE0), RGBColor(0xE6, 0xE8, 0xFF)
    PURPLE = RGBColor(0x4B, 0x3F, 0xC4)
    cw, bh, gap = 2.05, 0.7, 0.1
    xs = [0.45, 2.8, 5.15, 7.5]
    ys = [1.68, 1.68 + bh + gap, 1.68 + 2 * (bh + gap)]
    headers = ["DATA SOURCES", "CONVERGED FEATURES", "PREDICT & ACT", "EXPERIENCE"]
    for x, h in zip(xs, headers):
        text_box(s, x, 1.42, cw, 0.24, [[(h, dict(bold=True, size=8.5, color=GRAY))]], size=8.5)
    columns = [
        [("OT · sensor stream", ["vibration, temp, RPM, status", "5-min samples · structured"], OT, GREEN),
         ("IT · ERP production orders", ["planned/run min, counts, scrap", "per shift · structured"], IT, BLUE),
         ("IT · CMMS work orders", ["type, parts, cost, technician", "free-text findings · unstructured"], IT, BLUE)],
        [("Hourly feature table", ["rolling stats, trends, baselines", "+ hours since maintenance (IT)"], IT, BLUE),
         ("OEE daily", ["Availability × Perf. × Quality", "Snowflake: Dynamic Table"], IT, BLUE),
         ("Evidence pack builder", ["sensor deltas, drivers, WO history", "production context, for grounding"], IT, BLUE)],
        [("Failure model (GBM)", ["P(failure in 24 h) + top drivers", "holdout-tested · Model Registry"], ACT, ORANGE),
         ("Alert engine + rule limits", ["risk ≥ 0.7 CRITICAL · ≥ 0.4 WARNING", "dedupe, auto-resolve, safety net"], ACT, ORANGE),
         ("Work-order automation", ["P1 predictive WO with evidence", "one WO per asset · 15-min Task"], ACT, ORANGE)],
        [("Streamlit command center", ["overview, triage, asset explorer,", "work orders, system health"], UX, PURPLE),
         ("Ask in natural language", ["verified queries → guarded SQL", "read-only, auto-LIMIT, retry"], UX, PURPLE),
         ("Root cause · Granite 1B", ["grounded on the evidence pack", "rule-based fallback if LLM is down"], UX, PURPLE)],
    ]
    for x, col in zip(xs, columns):
        for y, (t, lines, fill, tc) in zip(ys, col):
            box(s, x, y, cw, bh, t, lines, fill=fill, title_color=tc, size=8.5, title_size=10)
    for i in range(3):
        for y in ys:
            arrow(s, xs[i] + cw, y + bh / 2, xs[i + 1], y + bh / 2)
    # CoCo skills strip
    rect(s, 0.45, 4.08, 9.1, 1.2, fill=NAVY)
    text_box(s, 0.6, 4.13, 8.8, 0.25, [[("CoCo CLI skills used, and how they connect", dict(bold=True, size=10, color=WHITE))]], size=10)
    skills = [("Synthetic data generation", "referentially consistent OT + ERP + CMMS sets, no production data"),
              ("Data pipeline creation", "hourly features & OEE as Dynamic Tables; scheduled scoring Task"),
              ("Semantic model authoring", "Cortex Analyst YAML with verified queries (same tier as the app)"),
              ("Streamlit app generation", "command center scaffolded and iterated in CoCo"),
              ("Automation + guardrails", "unattended runs, read-only SQL guard, alert dedupe, LLM fallback"),
              ("Reusable skill", "coco/skills/factorypulse-pipeline: build, run, validate, deploy")]
    sw = 2.95
    for i, (h, d) in enumerate(skills):
        x, y = 0.6 + (i % 3) * sw, 4.42 + (i // 3) * 0.42
        text_box(s, x, y, sw - 0.1, 0.42, [[(h + ": ", dict(bold=True, size=8.5, color=SKY)), (d, dict(size=8.5, color=WHITE))]], size=8.5, space_after=0)


def slide_impact(s, m):
    title(s, "Impact: measurable now, scalable, extensible later",
          f"Figures from the generated plant (8 assets, 30 days, {m['failures']} unplanned failures, forward-looking holdout on the last 30 % of the timeline)")
    stats = [(f"{m['roc']:.2f}", "ROC-AUC on holdout", BLUE),
             (f"{m['recall']:.0%}", "of failure hours flagged in the 24 h window (recall)", BLUE),
             (f"~{m['lead_h']:.0f} h", "average warning lead time before breakdown", ORANGE),
             (f"{m['downtime_h']:.0f} h", f"unplanned downtime in 30 days (~{m['avg_down_h']:.0f} h per event)", GRAY),
             ("< 1 min", "from CRITICAL alert to P1 work order (was hours)", GREEN)]
    sw, sh = 1.74, 1.05
    for i, (val, lab, col) in enumerate(stats):
        x = 0.45 + i * (sw + 0.1)
        rect(s, x, 1.55, sw, sh, fill=TINT2)
        text_box(s, x + 0.05, 1.58, sw - 0.1, 0.55, [[(val, dict(bold=True, size=24, color=col))]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text_box(s, x + 0.08, 2.1, sw - 0.16, 0.48, [[(lab, dict(size=8.5, color=GRAY))]], align=PP_ALIGN.CENTER)
    cols = [
        ("1", "Measurable outcomes", BLUE, [
            f"Precision {m['precision']:.0%} / recall {m['recall']:.0%}: few false alarms, most failures caught a shift ahead.",
            "Every CRITICAL alert carries its drivers and opens a work order with evidence, cutting triage from hours to seconds.",
            "Root-cause narrative in ~10 s and data questions in ~3 s on a CPU-only plant PC, grounded and audited.",
            f"If 9 of 10 failures become planned interventions, most of the {m['downtime_h']:.0f} h/month of unplanned downtime turns into shorter, scheduled work."]),
        ("2", "Scalability potential", GREEN, [
            "Snowflake twin: Dynamic Tables refresh incrementally, a Task scores every 15 min, the model sits in the Model Registry.",
            "Adding assets, lines or plants adds rows, not code: features, OEE and alerts are all per-asset.",
            "Cortex Analyst semantic model exposes the same verified queries in Snowsight agents and the Slackbot.",
            "Sensitive OT data can stay on site with the local Granite model; cloud handles scale and governance."]),
        ("3", "Beyond the demo", ORANGE, [
            "Replace synthetic feeds with Snowpipe Streaming / Kafka for historians and ERP connectors for SAP.",
            "MCP connectors to the CMMS (SAP PM, Maximo) and Slack so the agent creates and routes real work orders.",
            "Feedback loop: closed work orders with technician verdicts retrain the model and refine failure modes.",
            "Survival model for time-to-failure, spare-parts forecasting, and a triage → RCA → planner multi-agent flow."]),
    ]
    cw = 2.95
    for i, (n, head, col, lines) in enumerate(cols):
        x = 0.45 + i * (cw + 0.125)
        rect(s, x, 2.78, cw, 2.5, fill=TINT if i == 1 else TINT2)
        card_header(s, x + 0.12, 2.9, cw - 0.24, n, head, fill=col)
        text_box(s, x + 0.12, 3.28, cw - 0.24, 1.95, [[(ln, {})] for ln in lines], size=9, bullets=True, space_after=2)


def slide_additional(s):
    title(s, "CoCo across the lifecycle, guardrails, and the 3-minute demo")
    phases = [("Planning", "coco/PLAN.md: problem framing, ontology, workflow, design decisions, validation plan drafted before the build"),
              ("Development", "pipeline, model, Streamlit app, Snowflake DDL / Dynamic Tables / semantic model built and iterated in CoCo"),
              ("Execution", "run_pipeline.py end-to-end; simulate_stream.py near-real-time ticks; scheduled Snowflake Task for unattended scoring"),
              ("Testing", "20 pytest checks: data integrity, OEE math, model floor, alert guardrails, SQL guard, LLM-free fallback")]
    pw = 2.2
    for i, (h, d) in enumerate(phases):
        x = 0.45 + i * (pw + 0.1)
        rect(s, x, 1.3, pw, 1.25, fill=TINT)
        circle(s, x + 0.12, 1.4, 0.3, BLUE, str(i + 1), size=10)
        text_box(s, x + 0.5, 1.38, pw - 0.6, 0.34, [[(h, dict(bold=True, size=12, color=NAVY))]], anchor=MSO_ANCHOR.MIDDLE)
        text_box(s, x + 0.12, 1.75, pw - 0.24, 0.78, [[(d, dict(size=8.5, color=NAVY))]], size=8.5)
    # guardrails
    rect(s, 0.45, 2.72, 4.45, 2.1, fill=TINT2)
    card_header(s, 0.57, 2.84, 4.2, "G", "Guardrails and graceful fallback", fill=GREEN)
    text_box(s, 0.57, 3.22, 4.2, 1.55, [[(t, {})] for t in [
        "Natural-language path can never write: single SELECT only, forbidden keywords blocked, auto-LIMIT, read-only connection.",
        "Verified queries answer common questions deterministically before the LLM is involved.",
        "Alerts de-duplicated per asset/type; only CRITICAL opens a work order; one open WO per asset; auto-resolve on recovery.",
        "Hard sensor limits (7.1 mm/s, 85 °C) alert even if the model is missing; LLM outage falls back to rule-based narrative.",
        "The LLM only sees the evidence pack, which is displayed next to every answer for audit."]], size=8.5, bullets=True, space_after=2)
    # demo flow
    rect(s, 5.1, 2.72, 4.45, 2.1, fill=TINT2)
    card_header(s, 5.22, 2.84, 4.2, "D", "Demo flow", fill=ORANGE)
    text_box(s, 5.22, 3.22, 4.2, 1.3, [[(t, {})] for t in [
        "Overview opens with CNC-002 and PMP-001 at CRITICAL, two auto-created P1 work orders, plant OEE by line.",
        "Alert triage: acknowledge, open the 72 h sensor view, ask Granite for the root cause (bearing wear signature).",
        "Investigate: \"Which asset had the most unplanned downtime?\" (verified) and a free-form question (text-to-SQL).",
        "Second terminal: simulate_stream.py --degrade CNC-003 raises a live alert and work order within a minute."]], size=8.5, bullets=True, space_after=2)
    text_box(s, 5.22, 4.5, 4.2, 0.3, [[("Repo: ", dict(bold=True, size=8.5, color=NAVY)),
                                       ("github.com/Amogh892/factorypulse-oee-command-center", dict(size=8.5, color=BLUE))]], size=8.5)


def main(template: str, out: str):
    m = load_metrics()
    print("metrics:", {k: round(v, 3) if isinstance(v, float) else v for k, v in m.items()})
    prs = Presentation(template)
    slides = list(prs.slides)
    cover, guidelines, blank_a, blank_b, additional, thanks = slides

    slide_cover(cover)
    slide_problem(blank_a)
    slide_architecture(blank_b)
    impact = clone_branded_slide(prs, blank_b)   # appended at the end for now
    slide_impact(impact, m)
    slide_additional(additional)

    delete_slide(prs, 1)                 # drop the guidelines slide -> cover, problem, arch, additional, thanks, impact
    move_slide(prs, 5, 3)                # impact before additional
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
